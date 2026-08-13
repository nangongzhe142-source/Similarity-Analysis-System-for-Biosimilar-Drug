# -*- coding: utf-8 -*-
"""
Generate structured TypeScript data files for the biosimilar similarity
framework website from the source Excel workbook (single source of truth).

- Sheet "2.特性鉴定" -> src/data/characterization-items.ts (61 items)
  * Rows 9-10 share merged content (two items: 修饰1 / 修饰2).
  * Rows 59-60 are one merged item (HDX-MS + methyl NMR as two methods).
- Sheet "1.法规框架"  -> src/data/regulatory-framework.ts

`zh` text comes verbatim from the Excel cells (openpyxl, data_only=True).
`en` text comes from en_translations.py (machine-translation placeholders,
TODO: 校对英文).
"""
import json
import os
import re
import sys

import openpyxl

from en_translations import FIELD_EN, METHOD_EN, REGULATORY_EN

EXCEL_PATH = r"d:\生物类似药判别系统\生物类似药评价指导原则\V0.1生物类似药药学比对研究质量属性、检测方法及相似性评价原则汇总表(1).xlsx"
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_ITEMS = os.path.join(PROJECT_ROOT, "src", "data", "characterization-items.ts")
OUT_REGULATORY = os.path.join(PROJECT_ROOT, "src", "data", "regulatory-framework.ts")

CHARACTERIZATION_SHEET = "2.特性鉴定"
REGULATORY_SHEET = "1.法规框架"

# Column indices (1-based) of sheet "2.特性鉴定"
COL_GUIDELINE_TERM = 1
COL_ITEM_NAME = 2
COL_APPLICABILITY = 3
COL_PURPOSE = 4
COL_PRIMARY_METHOD = 5
COL_ORTHOGONAL_METHOD = 6
COL_DETECTION_INDICATORS = 7
COL_SIMILARITY_METHOD = 8
COL_JUDGING_PRINCIPLE = 9
COL_NUMERIC_LIMIT = 10
COL_REMARK = 11

FIELD_COLUMNS = {
    "guidelineTerm": COL_GUIDELINE_TERM,
    "itemName": COL_ITEM_NAME,
    "applicability": COL_APPLICABILITY,
    "purpose": COL_PURPOSE,
    "detectionIndicators": COL_DETECTION_INDICATORS,
    "similarityMethod": COL_SIMILARITY_METHOD,
    "judgingPrinciple": COL_JUDGING_PRINCIPLE,
    "numericLimit": COL_NUMERIC_LIMIT,
    "remark": COL_REMARK,
}

# (excel_row, item_id, category_key); array order == page display order.
ROW_CONFIG: list[tuple[int, str, str]] = [
    (2, "intact-mass", "primary-structure"),
    (3, "deglycosylated-intact-mass", "primary-structure"),
    (4, "light-chain-mass", "primary-structure"),
    (5, "non-deglycosylated-heavy-chain-mass", "primary-structure"),
    (6, "deglycosylated-heavy-chain-mass", "primary-structure"),
    (7, "ms1-sequence-coverage", "primary-structure"),
    (8, "msms-sequence-coverage", "primary-structure"),
    (9, "ptm-modification-1", "ptm-glycosylation"),
    (10, "ptm-modification-2", "ptm-glycosylation"),  # merged with row 9 (cols B-K)
    (11, "cdr-signature-peptides", "primary-structure"),
    (12, "n-c-terminal-sequence", "primary-structure"),
    (13, "free-thiol", "primary-structure"),
    (14, "disulfide-bonds", "primary-structure"),
    (15, "far-uv-cd", "higher-order-structure"),
    (16, "near-uv-cd", "higher-order-structure"),
    (17, "intrinsic-fluorescence", "higher-order-structure"),
    (18, "thermal-stability", "higher-order-structure"),
    (19, "other-higher-order-structure-methods", "higher-order-structure"),
    (20, "n-glycosylation-site-occupancy", "ptm-glycosylation"),
    (21, "glycan-g0f", "ptm-glycosylation"),
    (22, "glycan-g0", "ptm-glycosylation"),
    (23, "other-n-glycans", "ptm-glycosylation"),
    (24, "sialic-acid-ngna", "ptm-glycosylation"),
    (25, "sialic-acid-nana", "ptm-glycosylation"),
    (26, "molar-extinction-coefficient", "physicochemical"),
    (27, "isoelectric-point", "physicochemical"),
    (28, "sec-hmw-aggregates", "purity-size-variants"),
    (29, "sec-main-peak-monomer", "purity-size-variants"),
    (30, "sec-lmw-fragments", "purity-size-variants"),
    (31, "reduced-ce-sds-hc-lc-purity", "purity-size-variants"),
    (32, "reduced-ce-sds-fragments", "purity-size-variants"),
    (33, "non-reduced-ce-sds-main-peak", "purity-size-variants"),
    (34, "non-reduced-ce-sds-fragments", "purity-size-variants"),
    (35, "acidic-charge-variants", "charge-variants"),
    (36, "main-charge-peak", "charge-variants"),
    (37, "basic-charge-variants", "charge-variants"),
    (38, "target-binding-activity", "binding-bioactivity"),
    (39, "moa-related-bioactivity", "binding-bioactivity"),
    (40, "fcgri-cd64-binding", "binding-bioactivity"),
    (41, "fcgriia-cd32a-binding", "binding-bioactivity"),
    (42, "fcgriib-cd32b-binding", "binding-bioactivity"),
    (43, "fcgriiia-cd16a-binding", "binding-bioactivity"),
    (44, "fcrn-binding", "binding-bioactivity"),
    (45, "c1q-binding", "binding-bioactivity"),
    (46, "adcc", "binding-bioactivity"),
    (47, "cdc", "binding-bioactivity"),
    (48, "other-moa-related-functions", "binding-bioactivity"),
    (49, "other-product-related-impurities", "process-product-impurities"),
    (50, "protein-a-residual", "process-product-impurities"),
    (51, "residual-dna", "process-product-impurities"),
    (52, "residual-hcp", "process-product-impurities"),
    (53, "other-process-related-impurities", "process-product-impurities"),
    (54, "methionine-tryptophan-oxidation", "ptm-glycosylation"),
    (55, "asn-deamidation", "ptm-glycosylation"),
    (56, "n-terminal-pyroglutamate", "ptm-glycosylation"),
    (57, "c-terminal-lysine-processing", "ptm-glycosylation"),
    (58, "ft-ir-secondary-structure", "higher-order-structure"),
    (59, "hdx-ms-nmr-high-resolution", "higher-order-structure"),  # merged with row 60
    (61, "galactosylation-g1f-g2f", "ptm-glycosylation"),
    (62, "high-mannose-glycans", "ptm-glycosylation"),
    (63, "core-fucosylation", "ptm-glycosylation"),
]

MERGED_CONTENT_ROW = {10: 9}  # row 10 reads cols B-K from row 9

# Rows whose 表征项目 cell is "见表末补充": display the 指南原词 (e.g.
# "翻译后修饰—修饰1") as the item title and demote "见表末补充" to the
# small guideline-term line, per UI review feedback.
NAME_SWAPPED_ROWS = {9, 10}
HDX_ROW = 59
HDX_EXTRA_ROW = 60  # extra primary method (methyl NMR) + extra detection indicators

METHOD_SPLIT_PATTERN = re.compile(r"[；;\n]+")

missing_translations: list[str] = []


def cell_text(sheet, row_number: int, column_number: int) -> str:
    value = sheet.cell(row=row_number, column=column_number).value
    return str(value).strip() if value is not None else ""


def english_for_field(row_number: int, field_name: str) -> str:
    key = (row_number, field_name)
    if key in FIELD_EN:
        return FIELD_EN[key]
    missing_translations.append(f"field {key}")
    return "TODO: English translation pending"


def split_method_segments(raw_text: str) -> list[str]:
    segments = [segment.strip() for segment in METHOD_SPLIT_PATTERN.split(raw_text)]
    return [segment for segment in segments if segment and segment != "-"]


def english_for_method(row_number: int, kind: str, index: int) -> str:
    key = (row_number, kind, index)
    if key in METHOD_EN:
        return METHOD_EN[key]
    missing_translations.append(f"method {key}")
    return "TODO: English translation pending"


def build_methods(sheet, item_id: str, source_row: int) -> list[dict]:
    methods: list[dict] = []
    specs = [
        ("primary", COL_PRIMARY_METHOD, "p"),
        ("orthogonal", COL_ORTHOGONAL_METHOD, "o"),
    ]
    for method_type, column_number, kind in specs:
        source_rows = [source_row]
        if source_row == HDX_ROW and method_type == "primary":
            source_rows.append(HDX_EXTRA_ROW)
        sequence_number = 0
        for actual_row in source_rows:
            raw_text = cell_text(sheet, actual_row, column_number)
            if not raw_text or raw_text == "-":
                continue
            segments = split_method_segments(raw_text)
            for segment_index, segment in enumerate(segments):
                sequence_number += 1
                methods.append({
                    "id": f"{item_id}-{method_type}-{sequence_number}",
                    "name": {
                        "zh": segment,
                        "en": english_for_method(actual_row, kind, segment_index),
                    },
                    "type": method_type,
                    "contentPlaceholder": True,
                    "rawSourceText": {
                        "zh": raw_text,
                        "en": "; ".join(
                            english_for_method(actual_row, kind, i)
                            for i in range(len(segments))
                        ),
                    },
                })
    return methods


def build_characterization_items(workbook) -> list[dict]:
    sheet = workbook[CHARACTERIZATION_SHEET]
    items: list[dict] = []

    for excel_row, item_id, category_key in ROW_CONFIG:
        content_row = MERGED_CONTENT_ROW.get(excel_row, excel_row)

        def field(field_name: str, zh_row: int) -> dict:
            zh_text = cell_text(sheet, zh_row, FIELD_COLUMNS[field_name])
            return {"zh": zh_text, "en": english_for_field(zh_row, field_name)}

        # guidelineTerm always comes from the item's own row (differs for 修饰1/修饰2)
        guideline_term = field("guidelineTerm", excel_row)
        item_name = field("itemName", content_row)
        if excel_row in NAME_SWAPPED_ROWS:
            guideline_term, item_name = item_name, guideline_term
        applicability = field("applicability", content_row)
        purpose = field("purpose", content_row)
        detection_indicators = field("detectionIndicators", content_row)
        similarity_method = field("similarityMethod", content_row)
        judging_principle = field("judgingPrinciple", content_row)
        numeric_limit = field("numericLimit", content_row)
        remark = field("remark", content_row)

        if excel_row == HDX_ROW:
            extra = field("detectionIndicators", HDX_EXTRA_ROW)
            detection_indicators = {
                "zh": detection_indicators["zh"] + "\n" + extra["zh"],
                "en": detection_indicators["en"] + "\n" + extra["en"],
            }

        # Methods are parsed from the row that actually holds the merged content.
        items.append({
            "id": item_id,
            "category": category_key,
            "isSupplementary": applicability["zh"].startswith("补充项"),
            "guidelineTerm": guideline_term,
            "itemName": item_name,
            "applicability": applicability,
            "purpose": purpose,
            "detectionIndicators": detection_indicators,
            "similarityMethod": similarity_method,
            "judgingPrinciple": judging_principle,
            "numericLimit": numeric_limit,
            "remark": remark,
            "methods": build_methods(sheet, item_id, content_row),
            "analysisPlaceholder": {
                "candidateDataSlot": True,
                "referenceDataSlot": True,
                "resultSlot": True,
            },
        })
    return items


RELATION_MAP = {
    "直接相关": "directly-related",
    "间接相关": "indirectly-related",
    "支持性": "supportive",
}


def english_for_regulatory(row_number: int, field_name: str, zh_text: str) -> str:
    key = (row_number, field_name)
    if key in REGULATORY_EN:
        return REGULATORY_EN[key]
    missing_translations.append(f"regulatory {key} ({zh_text[:20]})")
    return "TODO: English translation pending"


def build_regulatory_framework(workbook) -> dict:
    sheet = workbook[REGULATORY_SHEET]

    title_zh = cell_text(sheet, 1, 1)
    requirements: list[dict] = []
    last_ctd_section = ""
    for row_number in range(3, 12):  # rows 3-11
        ctd_section = cell_text(sheet, row_number, 1) or last_ctd_section
        last_ctd_section = ctd_section
        subject_zh = cell_text(sheet, row_number, 2)
        requirement_zh = cell_text(sheet, row_number, 3)
        page_reference = cell_text(sheet, row_number, 4)
        remark_zh = cell_text(sheet, row_number, 5)
        remark_zh_normalized = "" if remark_zh == "-" else remark_zh
        remark_en = english_for_regulatory(row_number, "remark", remark_zh)
        requirements.append({
            "id": f"requirement-{row_number - 2}",
            "ctdSection": ctd_section,
            "subject": {
                "zh": subject_zh,
                "en": english_for_regulatory(row_number, "subject", subject_zh),
            },
            "requirement": {
                "zh": requirement_zh,
                "en": english_for_regulatory(row_number, "requirement", requirement_zh),
            },
            "pageReference": page_reference,
            "remark": {
                "zh": remark_zh_normalized,
                "en": "" if remark_en == "-" else remark_en,
            },
        })

    relations: list[dict] = []
    for row_number in range(15, 24):  # rows 15-23
        ctd_section = cell_text(sheet, row_number, 1)
        subject_zh = cell_text(sheet, row_number, 2)
        relation_zh = cell_text(sheet, row_number, 3)
        relation_key = RELATION_MAP.get(relation_zh)
        if relation_key is None:
            print(f"WARNING: unknown relation value '{relation_zh}' at row {row_number}")
            relation_key = "supportive"
        relations.append({
            "id": f"relation-{row_number - 14}",
            "ctdSection": ctd_section,
            "subject": {
                "zh": subject_zh,
                "en": english_for_regulatory(row_number, "subject", subject_zh),
            },
            "relation": relation_key,
        })

    return {
        "sourceTitle": {
            "zh": title_zh,
            "en": english_for_regulatory(1, "title", title_zh),
        },
        "requirements": requirements,
        "relations": relations,
    }


HEADER_TEMPLATE = """\
// AUTO-GENERATED FILE — do not edit by hand.
// Source of truth: 生物类似药评价指导原则/V0.1生物类似药药学比对研究质量属性、检测方法及相似性评价原则汇总表(1).xlsx
//   sheet: {sheet_name} (read via openpyxl with data_only=True)
// Regenerate with: python scripts/generate_data.py
// NOTE: all `en` strings are machine-translation placeholders.
// TODO: 校对英文 (review the English translations).
"""


def write_typescript(path: str, sheet_name: str, import_line: str, declaration: str, payload) -> None:
    body = json.dumps(payload, ensure_ascii=False, indent=2)
    content = (
        HEADER_TEMPLATE.format(sheet_name=sheet_name)
        + import_line + "\n\n"
        + declaration + " = " + body + ";\n"
    )
    with open(path, "w", encoding="utf-8", newline="\n") as file:
        file.write(content)


def main() -> None:
    workbook = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    items = build_characterization_items(workbook)
    regulatory = build_regulatory_framework(workbook)

    assert len(items) == 61, f"Expected 61 items, got {len(items)}"
    supplementary_count = sum(1 for item in items if item["isSupplementary"])
    assert supplementary_count == 9, f"Expected 9 supplementary items, got {supplementary_count}"
    assert len(regulatory["requirements"]) == 9
    assert len(regulatory["relations"]) == 9

    category_counts: dict[str, int] = {}
    for item in items:
        category_counts[item["category"]] = category_counts.get(item["category"], 0) + 1

    write_typescript(
        OUT_ITEMS,
        CHARACTERIZATION_SHEET,
        'import type { CharacterizationItem } from "@/types/models";',
        "export const characterizationItems: CharacterizationItem[]",
        items,
    )
    write_typescript(
        OUT_REGULATORY,
        REGULATORY_SHEET,
        'import type { RegulatoryFramework } from "@/types/models";',
        "export const regulatoryFramework: RegulatoryFramework",
        regulatory,
    )

    print(f"Items: {len(items)} (supplementary: {supplementary_count})")
    print("Per category:", json.dumps(category_counts, ensure_ascii=False, indent=2))
    total_methods = sum(len(item["methods"]) for item in items)
    print(f"Total detection methods: {total_methods}")
    print(f"Regulatory requirements: {len(regulatory['requirements'])}, relations: {len(regulatory['relations'])}")
    if missing_translations:
        print(f"\nMISSING TRANSLATIONS ({len(missing_translations)}):")
        for entry in missing_translations:
            print("  -", entry)
        sys.exit(1)
    print("All translations resolved. Data files written.")


if __name__ == "__main__":
    main()
