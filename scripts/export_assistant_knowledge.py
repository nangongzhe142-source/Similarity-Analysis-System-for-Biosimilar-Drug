# -*- coding: utf-8 -*-
"""Export traceable Markdown for the Dify Chatflow knowledge base.

Reads the V2 workbook, site sidecars, guideline files and README files.
Does not read node_modules, build output, logs or mass-spec fixtures.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import sys
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKSPACE_ROOT = PROJECT_ROOT.parent
GUIDELINE_DIR = WORKSPACE_ROOT / "生物类似药评价指导原则"
OUT_DIR = PROJECT_ROOT / "knowledge" / "assistant"
SOURCE_WORKBOOK_NAME = (
    "V2-生物类似药药学比对研究质量属性、检测方法及相似性评价原则汇总表.xlsx"
)
EXCEL_PATH = GUIDELINE_DIR / SOURCE_WORKBOOK_NAME
EXPECTED_EXCEL_SHA256 = (
    "8bd6b18f08d3b9a3b99c61cdb8eeff24ceef75a54a43b9f7d14d1709b703051f"
)

GUIDELINE_NATURE = {
    "《生物类似药相似性评价和适应症外推技术指导原则》.pdf": "真实监管资料（指导原则）",
    "生物类似药研发与评价技术指导原则（试行）.doc": "真实监管资料（指导原则）",
    "生物类似药首次申报临床试验药学资料撰写指导原则.pdf": "真实监管资料（指导原则）",
    "生物类似药药学评价比较.docx": "方法学资料（药学评价比较说明）",
    "判断相似性的标准.docx": "方法学资料（待与正式指导原则对照）",
    "V3.7-生物类似药课题中国部分-20260629韩欣烘.docx": "课题内部稿（非正式监管文件）",
    "20260817组会会议记录.docx": "内部会议记录（非监管结论）",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.replace("\r\n", "\n").strip() + "\n", encoding="utf-8")


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_ts_json_array(path: Path) -> list:
    source = path.read_text(encoding="utf-8")
    start = source.index("= [") + 2
    end = source.rindex("]") + 1
    return json.loads(source[start:end])


def load_ts_json_object(path: Path) -> dict:
    source = path.read_text(encoding="utf-8")
    start = source.index("= {")
    start = source.index("{", start)
    end = source.rindex("}") + 1
    return json.loads(source[start:end])


def unescape_ts_string(raw: str) -> str:
    return raw.replace(r"\"", '"').replace(r"\n", "\n").replace(r"\\", "\\")


def extract_quoted(pattern: str, block: str) -> str:
    match = re.search(pattern, block)
    return unescape_ts_string(match.group(1)) if match else ""


def parse_localized(block: str, field: str) -> dict[str, str]:
    match = re.search(
        rf"{field}:\s*\{{\s*zh:\s*\"((?:[^\"\\]|\\.)*)\"\s*,\s*en:\s*\"((?:[^\"\\]|\\.)*)\"",
        block,
        re.S,
    )
    if not match:
        zh = extract_quoted(rf"{field}:\s*\{{[\s\S]*?zh:\s*\"((?:[^\"\\]|\\.)*)\"", block)
        en = extract_quoted(rf"{field}:\s*\{{[\s\S]*?en:\s*\"((?:[^\"\\]|\\.)*)\"", block)
        return {"zh": zh, "en": en}
    return {
        "zh": unescape_ts_string(match.group(1)),
        "en": unescape_ts_string(match.group(2)),
    }


def parse_assign_method_content(path: Path) -> list[dict]:
    source = path.read_text(encoding="utf-8")
    rows: list[dict] = []
    for match in re.finditer(
        r"assign\(\s*byMethodId,\s*(\[[^\]]+\])\s*,\s*\{(.*?)\}\s*,\s*\)",
        source,
        re.S,
    ):
        method_ids = re.findall(r"\"([a-z0-9-]+)\"", match.group(1))
        body = match.group(2)
        zh_match = re.search(r"zh:\s*\"((?:[^\"\\]|\\.)*)\"", body, re.S)
        en_match = re.search(r"en:\s*\"((?:[^\"\\]|\\.)*)\"", body, re.S)
        principle_zh = unescape_ts_string(zh_match.group(1)) if zh_match else ""
        principle_en = unescape_ts_string(en_match.group(1)) if en_match else ""
        for method_id in method_ids:
            rows.append(
                {
                    "methodId": method_id,
                    "principleZh": principle_zh,
                    "principleEn": principle_en,
                }
            )
    return rows


def parse_scheme_blocks(path: Path) -> list[dict]:
    source = path.read_text(encoding="utf-8")
    starts = [
        (match.start(), match.group(1))
        for match in re.finditer(r"^\s+itemId:\s*\"([a-z0-9-]+)\"", source, re.M)
    ]
    blocks = []
    for index, (offset, item_id) in enumerate(starts):
        end = starts[index + 1][0] if index + 1 < len(starts) else len(source)
        block = source[offset:end]
        method_ids = re.findall(
            r"\"([a-z0-9-]+)\"",
            extract_quoted(r"methodIds:\s*\[([^\]]*)\]", block),
        )
        source_row = extract_quoted(r"provenance\(\s*(\d+)", block)
        completeness = extract_quoted(r"completeness:\s*\"([a-z-]+)\"", block)
        blocks.append(
            {
                "itemId": item_id,
                "methodIds": method_ids,
                "completeness": completeness,
                "sourceRow": source_row,
                "guidelineTerm": parse_localized(block, "guidelineTerm"),
                "notDefinedReason": parse_localized(block, "notDefinedReason"),
                "recognizedContent": parse_localized(block, "recognizedContent"),
                "comparisonBaseline": parse_localized(block, "comparisonBaseline"),
                "decisionMethod": parse_localized(block, "decisionMethod"),
                "numericBoundary": parse_localized(block, "numericBoundary"),
                "finalProgramRule": parse_localized(block, "finalProgramRule"),
                "basis": parse_localized(block, "basis"),
            }
        )
    return blocks


def extract_case_blocks(path: Path) -> list[dict]:
    source = path.read_text(encoding="utf-8")
    cases: list[dict] = []
    for match in re.finditer(r"^\s{4,6}id: \"([^\"]+)\",$", source, re.M):
        case_id = match.group(1)
        start = match.start()
        depth = 0
        end = len(source)
        for index, character in enumerate(source[start:], start):
            if character == "{":
                depth += 1
            elif character == "}":
                if depth == 0:
                    end = index
                    break
                depth -= 1
        block = source[start:end]
        item_match = re.search(
            r"^\s{2}\"([a-z0-9-]+)\": \[$", source[:start], re.M
        )
        citation_call = re.search(
            r"gp2015Source\(\s*\"((?:[^\"\\]|\\.)*)\"\s*,\s*\"((?:[^\"\\]|\\.)*)\"",
            block,
            re.S,
        )
        cases.append(
            {
                "itemId": item_match.group(1) if item_match else "",
                "caseId": case_id,
                "evidenceLevel": extract_quoted(r"evidenceLevel: \"([^\"]+)\"", block),
                "headlineZh": extract_quoted(r"headline:\s*\{\s*zh:\s*\"((?:[^\"\\]|\\.)*)\"", block),
                "citationZh": citation_call.group(1) if citation_call else "",
                "citationEn": citation_call.group(2) if citation_call else "",
                "hasSource": "gp2015Source(" in block,
                "sourceFile": path.name,
            }
        )
    return cases


def parse_analysis_config(path: Path) -> list[dict]:
    source = path.read_text(encoding="utf-8")
    rows = []
    for match in re.finditer(
        r"methodId:\s*\"([a-z0-9-]+)\"[\s\S]{0,400}?status:\s*\"([a-z-]+)\"",
        source,
    ):
        rows.append({"methodId": match.group(1), "status": match.group(2)})
    return rows


def export_excel(excel_path: Path, out_dir: Path, workbook_sha: str) -> list[Path]:
    import openpyxl

    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    written: list[Path] = []
    for sheet in workbook.worksheets:
        lines = [
            f"# {sheet.title}",
            "",
            f"- sourceFile: {SOURCE_WORKBOOK_NAME}",
            f"- sourceSheet: {sheet.title}",
            f"- sourceWorkbookSha256: {workbook_sha}",
            f"- nature: 系统规则（V2 汇总表原文）",
            "",
        ]
        for row_index, row in enumerate(sheet.iter_rows(values_only=False), start=1):
            nonempty = [cell for cell in row if cell.value not in (None, "")]
            if not nonempty:
                continue
            lines.append(f"## 行 {row_index}")
            for cell in nonempty:
                lines.append(
                    f"- {sheet.title}!{cell.coordinate}: {cell_text(cell.value)}"
                )
            lines.append("")
        target = out_dir / "01-v2-excel" / f"{sheet.title}.md"
        write_text(target, "\n".join(lines))
        written.append(target)
    return written


def extract_pdf(path: Path) -> str:
    import fitz

    document = fitz.open(path)
    pages = []
    for index, page in enumerate(document, start=1):
        pages.append(f"## 页 {index}\n\n{page.get_text().strip()}")
    document.close()
    return "\n\n".join(pages)


def extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        chunks.append(f"## 表 {table_index}")
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            chunks.append(f"- 行{row_index}: " + " | ".join(cells))
    return "\n\n".join(chunks)


def extract_doc(path: Path) -> str:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    document = word.Documents.Open(str(path), ReadOnly=True)
    try:
        return document.Content.Text
    finally:
        document.Close(False)
        word.Quit()
        pythoncom.CoUninitialize()


def export_guidelines(out_dir: Path) -> list[Path]:
    written: list[Path] = []
    for path in sorted(GUIDELINE_DIR.iterdir()):
        if path.name.startswith("~$") or path.is_dir():
            continue
        if path.suffix.lower() == ".xlsx":
            continue
        nature = GUIDELINE_NATURE.get(path.name, "未分类资料")
        header = [
            f"# {path.name}",
            "",
            f"- sourceFile: 生物类似药评价指导原则/{path.name}",
            f"- nature: {nature}",
            f"- byteSize: {path.stat().st_size}",
            f"- sha256: {sha256_file(path)}",
            "",
            "以下为抽取文本。页码或段落号来自抽取过程，不是后补编造。",
            "",
        ]
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                body = extract_pdf(path)
            elif suffix == ".docx":
                body = extract_docx(path)
            elif suffix == ".doc":
                body = extract_doc(path)
            else:
                body = f"（未抽取：后缀 {suffix}）"
        except Exception as error:  # noqa: BLE001
            body = f"（抽取失败：{type(error).__name__}: {error}）"
        target = out_dir / "02-guidelines" / f"{path.stem}.md"
        write_text(target, "\n".join(header) + body)
        written.append(target)
    return written


def export_catalog(out_dir: Path) -> Path:
    items = load_ts_json_array(PROJECT_ROOT / "src" / "data" / "characterization-items.ts")
    lines = [
        "# 检测项目与方法目录",
        "",
        "- sourceFile: src/data/characterization-items.ts",
        "- nature: 系统规则（网页实际展示数据，源自 V2 Excel）",
        "",
    ]
    for item in items:
        lines.append(f"## itemId `{item['id']}`")
        lines.append(f"- categoryKey: {item['category']}")
        lines.append(f"- itemName.zh: {item['itemName']['zh']}")
        lines.append(f"- itemName.en: {item['itemName']['en']}")
        lines.append(f"- guidelineTerm.zh: {item['guidelineTerm']['zh']}")
        lines.append(f"- isSupplementary: {item['isSupplementary']}")
        lines.append(f"- judgingPrinciple.zh: {item['judgingPrinciple']['zh']}")
        lines.append(f"- numericLimit.zh: {item['numericLimit']['zh']}")
        for method in item.get("methods", []):
            lines.append(
                f"- methodId `{method['id']}` type={method['type']} "
                f"name.zh={method['name']['zh']}"
            )
        lines.append("")
    target = out_dir / "03-site-data" / "characterization-catalog.md"
    write_text(target, "\n".join(lines))
    return target


def export_method_content(out_dir: Path, items: list) -> Path:
    content = parse_assign_method_content(PROJECT_ROOT / "src" / "data" / "method-content.ts")
    item_by_method = {}
    for item in items:
        for method in item.get("methods", []):
            item_by_method[method["id"]] = item
    lines = [
        "# 方法学正文（原理）",
        "",
        "- sourceFile: src/data/method-content.ts",
        "- nature: 方法学资料",
        "- note: 当前仅嵌入 principle；SOP 其余字段待嵌入。",
        "",
    ]
    for row in content:
        item = item_by_method.get(row["methodId"], {})
        lines.append(f"## methodId `{row['methodId']}`")
        lines.append(f"- itemId: {item.get('id', '')}")
        lines.append(f"- categoryKey: {item.get('category', '')}")
        lines.append(f"- language: zh/en")
        lines.append(f"- principle.zh: {row['principleZh']}")
        lines.append(f"- principle.en: {row['principleEn']}")
        lines.append("")
    target = out_dir / "03-site-data" / "method-content.md"
    write_text(target, "\n".join(lines))
    return target


def export_schemes(out_dir: Path) -> Path:
    schemes = parse_scheme_blocks(PROJECT_ROOT / "src" / "data" / "similarity-schemes.ts")
    lines = [
        "# V2 Sheet3 相似性评价规则 sidecar",
        "",
        f"- sourceFile: src/data/similarity-schemes.ts",
        f"- sourceWorkbook: {SOURCE_WORKBOOK_NAME}",
        f"- sourceWorkbookSha256: {EXPECTED_EXCEL_SHA256}",
        "- sourceSheet: 3.特性鉴定相似性评价方案",
        "- nature: 系统规则",
        "",
    ]
    for scheme in schemes:
        lines.append(f"## itemId `{scheme['itemId']}`")
        lines.append(f"- sourceRow: {scheme['sourceRow']}")
        lines.append(f"- completeness: {scheme['completeness']}")
        lines.append(f"- methodIds: {', '.join(scheme['methodIds'])}")
        lines.append(f"- guidelineTerm.zh: {scheme['guidelineTerm']['zh']}")
        if scheme["completeness"] != "complete":
            lines.append(
                f"- notDefinedReason.zh: {scheme['notDefinedReason']['zh']}"
            )
        else:
            lines.append(f"- recognizedContent.zh: {scheme['recognizedContent']['zh']}")
            lines.append(
                f"- comparisonBaseline.zh: {scheme['comparisonBaseline']['zh']}"
            )
            lines.append(f"- decisionMethod.zh: {scheme['decisionMethod']['zh']}")
            lines.append(f"- numericBoundary.zh: {scheme['numericBoundary']['zh']}")
            lines.append(f"- finalProgramRule.zh: {scheme['finalProgramRule']['zh']}")
            lines.append(f"- basis.zh: {scheme['basis']['zh']}")
        lines.append("")
    target = out_dir / "03-site-data" / "similarity-schemes.md"
    write_text(target, "\n".join(lines))
    return target


def export_regulatory(out_dir: Path) -> Path:
    data = load_ts_json_object(PROJECT_ROOT / "src" / "data" / "regulatory-framework.ts")
    lines = [
        "# 法规框架（网页展示）",
        "",
        "- sourceFile: src/data/regulatory-framework.ts",
        "- nature: 系统规则 / 法规对照",
        f"- sourceTitle.zh: {data['sourceTitle']['zh']}",
        "",
    ]
    for row in data["requirements"]:
        lines.append(f"## {row['id']}")
        lines.append(f"- ctdSection: {row['ctdSection']}")
        lines.append(f"- pageReference: {row['pageReference']}")
        lines.append(f"- subject.zh: {row['subject']['zh']}")
        lines.append(f"- requirement.zh: {row['requirement']['zh']}")
        lines.append("")
    for row in data["relations"]:
        lines.append(f"## {row['id']}")
        lines.append(f"- ctdSection: {row['ctdSection']}")
        lines.append(f"- relation: {row['relation']}")
        lines.append(f"- subject.zh: {row['subject']['zh']}")
        lines.append("")
    target = out_dir / "03-site-data" / "regulatory-framework.md"
    write_text(target, "\n".join(lines))
    return target


def export_analysis_config(out_dir: Path) -> Path:
    rows = parse_analysis_config(
        PROJECT_ROOT / "src" / "data" / "method-analysis-config.ts"
    )
    lines = [
        "# 方法分析面板状态",
        "",
        "- sourceFile: src/data/method-analysis-config.ts",
        "- nature: 系统操作说明",
        "",
    ]
    for row in rows:
        lines.append(f"- methodId `{row['methodId']}` status={row['status']}")
    target = out_dir / "03-site-data" / "method-analysis-status.md"
    write_text(target, "\n".join(lines))
    return target


def export_cases(out_dir: Path) -> tuple[Path, Path]:
    cases: list[dict] = []
    for name in ("reference-cases.ts", "reference-cases-gp2015-remaining.ts"):
        cases.extend(extract_case_blocks(PROJECT_ROOT / "src" / "data" / name))
    header_common = [
        "- sourceProduct: GP2015 / Erelzi / BLA 761042",
        "- regulator: FDA",
        "- localSourcePath: 生物类似药审批报告/翻译/output/17_etanercept_szzs/",
        "",
    ]
    regulatory = [
        "# 监管审评案例（已溯源）",
        "",
        "- nature: 真实监管资料（转录自公开审评报告，英文原文未逐条核对）",
        *header_common,
    ]
    illustrative = [
        "# 示意案例（非真实监管数据）",
        "",
        "- nature: 模拟数据 / 示意说明，不得当作监管结论",
        *header_common,
    ]
    for case in cases:
        body = [
            f"## caseId `{case['caseId']}`",
            f"- itemId: {case['itemId']}",
            f"- evidenceLevel: {case['evidenceLevel']}",
            f"- sourceFile: src/data/{case['sourceFile']}",
            f"- citation.zh: {case['citationZh']}",
            f"- headline.zh: {case['headlineZh']}",
            "",
        ]
        if case["evidenceLevel"] == "illustrative":
            illustrative.extend(body)
        else:
            regulatory.extend(body)
    reg_path = out_dir / "04-cases" / "regulatory-gp2015.md"
    ill_path = out_dir / "04-cases" / "illustrative.md"
    write_text(reg_path, "\n".join(regulatory))
    write_text(ill_path, "\n".join(illustrative))
    return reg_path, ill_path


def export_readmes(out_dir: Path) -> list[Path]:
    mapping = [
        (PROJECT_ROOT / "README.md", "system-readme.md", "系统操作说明"),
        (
            PROJECT_ROOT / "analysis-service" / "README.md",
            "analysis-service-readme.md",
            "分析服务操作说明",
        ),
    ]
    written = []
    for source, name, nature in mapping:
        text = source.read_text(encoding="utf-8")
        target = out_dir / "05-system" / name
        write_text(
            target,
            "\n".join(
                [
                    f"# {source.name}",
                    "",
                    f"- sourceFile: {source.relative_to(PROJECT_ROOT).as_posix()}",
                    f"- nature: {nature}",
                    "",
                    text,
                ]
            ),
        )
        written.append(target)
    return written


def main() -> int:
    if not EXCEL_PATH.is_file():
        print(f"FAIL  Excel not found: {EXCEL_PATH}", file=sys.stderr)
        return 1
    workbook_sha = sha256_file(EXCEL_PATH)
    if workbook_sha != EXPECTED_EXCEL_SHA256:
        print(
            f"FAIL  Excel SHA-256 {workbook_sha} != {EXPECTED_EXCEL_SHA256}",
            file=sys.stderr,
        )
        return 1

    if OUT_DIR.exists():
        for child in OUT_DIR.rglob("*"):
            if child.is_file():
                child.unlink()

    written: list[Path] = []
    written.extend(export_excel(EXCEL_PATH, OUT_DIR, workbook_sha))
    written.extend(export_guidelines(OUT_DIR))
    items = load_ts_json_array(PROJECT_ROOT / "src" / "data" / "characterization-items.ts")
    written.append(export_catalog(OUT_DIR))
    written.append(export_method_content(OUT_DIR, items))
    written.append(export_schemes(OUT_DIR))
    written.append(export_regulatory(OUT_DIR))
    written.append(export_analysis_config(OUT_DIR))
    written.extend(export_cases(OUT_DIR))
    written.extend(export_readmes(OUT_DIR))

    required_tokens = [
        "sourceSheet",
        "methodId",
        "itemId",
        "evidenceLevel",
        "nature",
    ]
    checks = {
        "sourceSheet": any("sourceSheet" in p.read_text(encoding="utf-8") for p in written),
        "methodId": (OUT_DIR / "03-site-data" / "method-content.md").read_text(
            encoding="utf-8"
        ).count("methodId")
        > 10,
        "itemId": (OUT_DIR / "03-site-data" / "characterization-catalog.md")
        .read_text(encoding="utf-8")
        .count("itemId")
        > 50,
        "evidenceLevel": "evidenceLevel" in (OUT_DIR / "04-cases" / "regulatory-gp2015.md").read_text(
            encoding="utf-8"
        ),
        "illustrative_separated": (OUT_DIR / "04-cases" / "illustrative.md").is_file(),
        "nature": all("nature:" in p.read_text(encoding="utf-8")[:800] for p in written),
    }
    failed = [name for name, ok in checks.items() if not ok]
    if failed:
        print(f"FAIL  provenance checks: {failed}", file=sys.stderr)
        return 1

    manifest_lines = [
        "# 助手知识导出清单",
        "",
        f"- exportedOn: {date.today().isoformat()}",
        f"- excelSha256: {workbook_sha}",
        "",
    ]
    for path in written:
        rel = path.relative_to(OUT_DIR).as_posix()
        manifest_lines.append(f"- {rel} ({path.stat().st_size} bytes)")
    write_text(OUT_DIR / "00-manifest.md", "\n".join(manifest_lines))
    print(f"OK  {len(written)} files -> {OUT_DIR}")
    for name in required_tokens:
        print(f"  token {name}: present")
    return 0


if __name__ == "__main__":
    sys.exit(main())
